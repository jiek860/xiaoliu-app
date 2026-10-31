# -*- coding: utf-8 -*-
"""文档解析：PDF / Word / Excel / PPT / Markdown / HTML / 文本。"""
import os
import re


def _clean(text):
    text = str(text or "").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _cloud_ocr(trunk, path):
    """用模块里的云端多模态 OCR 兜底，返回 {text} 或 {error}。"""
    if trunk is None:
        return {"error": "没有可用的云 OCR（未传入躯干）"}
    try:
        mod = (getattr(trunk, "modules", {}) or {}).get("image_ocr")
        if mod is None:
            return {"error": "云 OCR 模块未加载"}
        r = mod.recognize(path) or {}
        text = str(r.get("文字") or r.get("text") or "").strip()
        if text:
            return {"text": text}
        return {"error": str(r.get("错误") or r.get("error") or "云 OCR 没返回文字")[:120]}
    except Exception as e:
        return {"error": "云 OCR 失败: " + str(e)[:100]}


def parse(path, trunk=None):
    path = os.path.abspath(str(path or ""))
    if not os.path.isfile(path):
        return {"ok": False, "error": "文件不存在", "path": path}
    try:
        if os.path.getsize(path) > 50 * 1024 * 1024:
            return {"ok": False, "error": "文件超过 50MB，已拒绝解析",
                    "path": path}
    except OSError:
        pass
    ext = os.path.splitext(path)[1].lower()
    text = ""
    meta = {"title": os.path.basename(path), "source": path}
    try:
        if ext in (".txt", ".md", ".log", ".csv", ".json", ".yaml", ".yml"):
            text = open(path, encoding="utf-8", errors="replace").read()
        elif ext in (".html", ".htm"):
            raw = open(path, encoding="utf-8", errors="replace").read()
            text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", raw,
                          flags=re.S | re.I)
            text = re.sub(r"<[^>]+>", " ", text)
        elif ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            reader = PdfReader(path)
            pages = []
            for i, p in enumerate(reader.pages, 1):
                t = p.extract_text() or ""
                pages.append({"page": i, "text": _clean(t)})
            text = "\n\n".join(p["text"] for p in pages)
            meta["pages"] = pages
        elif ext == ".docx":
            from docx import Document
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
        elif ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append("## " + ws.title)
                for row in ws.iter_rows(values_only=True):
                    vals = [str(x) for x in row if x is not None]
                    if vals:
                        rows.append(" | ".join(vals))
            text = "\n".join(rows)
        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"):
            # 本地 OCR 优先（离线、免费）；本地缺依赖或没认出字，
            # 再用模块里的云端多模态 OCR 兜底（modules/installed/image_ocr）。
            local_err = ""
            try:
                from PIL import Image
                import pytesseract
            except ImportError:
                local_err = "本地 OCR 需要 Pillow + pytesseract（并安装 Tesseract）"
            else:
                try:
                    img = Image.open(path)
                    text = pytesseract.image_to_string(img, lang="chi_sim+eng")
                    meta["ocr"] = "local"
                except Exception as e:
                    local_err = "本地 OCR 失败: " + str(e)[:80]
            if not _clean(text):
                cloud = _cloud_ocr(trunk, path)
                if cloud.get("text"):
                    text = cloud["text"]
                    meta["ocr"] = "cloud"
                else:
                    reason = cloud.get("error") or local_err or "图片里没有文字"
                    return {"ok": False, "path": path, "skipped": True,
                            "error": "图片没识别出文字（%s）。可在「设置 → 依赖」"
                                     "装本地 OCR，或配置云端模型后用云 OCR。"
                                     % reason}
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                parts.append("## Slide %d" % i)
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
            text = "\n".join(parts)
        else:
            return {"ok": False, "error": "暂不支持的文件类型: " + ext,
                    "path": path}
    except ImportError as e:
        return {"ok": False, "error": "缺少文档解析依赖: " + str(e)[:80],
                "path": path}
    except Exception as e:
        return {"ok": False, "error": "解析失败: " + str(e)[:120],
                "path": path}
    text = _clean(text)
    if not text:
        # 扫描件 PDF / 空表格里常见：以前返回 ok=True + 空文本，
        # 附件入口只会显示"没有可提取的文字"，用户不知道该怎么办。
        return {"ok": False, "path": path, "skipped": True,
                "error": "文件里没有可提取的文字（扫描件可先转成图片走 OCR）"}
    return {"ok": True, "path": path, "text": text,
            "chars": len(text), "metadata": meta}


def chunk_text(text, chunk_size=800, overlap=120):
    text = _clean(text)
    if not text:
        return []
    chunks = []
    step = max(1, chunk_size - overlap)
    for i in range(0, len(text), step):
        piece = text[i:i + chunk_size].strip()
        if piece:
            chunks.append({"id": "chunk_%d" % (len(chunks) + 1),
                           "text": piece, "offset": i})
        if i + chunk_size >= len(text):
            break
    return chunks
